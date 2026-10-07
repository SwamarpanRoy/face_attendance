"""Admin UI (M7): live view, edits with reason and audit, roles, reports, devices, audit."""

from __future__ import annotations

import io
import re
import uuid
from datetime import UTC, datetime, timedelta

from openpyxl import load_workbook
from sqlalchemy import select

from server.app.models import (
    EMBEDDING_BYTES,
    Attendance,
    AttendanceMethod,
    AttendanceSession,
    AttendanceStatus,
    AuditLog,
    Consent,
    ConsentMethod,
    Course,
    CourseSection,
    Device,
    FaceTemplate,
    SessionStatus,
    TemplateSource,
)

ROW_RE = re.compile(
    r'<tr class="(short)?">\s*<td>22EC71</td>\s*<td><a[^>]*>(\w+)</a></td>.*?'
    r"<td>(\d+)</td><td>(\d+)</td><td>(\d+)</td><td>(\d+)</td>\s*<td>([^<]*)</td>",
    re.S,
)
TOKEN_RE = re.compile(r'<code id="token">([^<]+)</code>')
FORM_RE = re.compile(r"<form\b([^>]*)>(.*?)</form>", re.S | re.I)


def _session(db, course, section, faculty, device=None, *, started_at=None, ended=True):
    row = AttendanceSession(
        id=uuid.uuid4(),
        course_id=course.id,
        section_id=section.id,
        faculty_id=faculty.id,
        device_id=device.id if device else None,
        started_at=started_at or datetime.now(UTC) - timedelta(hours=1),
        ended_at=(datetime.now(UTC) - timedelta(minutes=5)) if ended else None,
        status=SessionStatus.ENDED if ended else SessionStatus.LIVE,
    )
    db.add(row)
    db.commit()
    return row


def _mark(db, session, student, status, *, method=AttendanceMethod.FACE, score=0.8):
    row = Attendance(
        session_id=session.id,
        student_id=student.id,
        status=status,
        method=method,
        score=score if method is AttendanceMethod.FACE else None,
        captured_at=datetime.now(UTC) if method is AttendanceMethod.FACE else None,
        event_uuid=uuid.uuid4() if method is AttendanceMethod.FACE else None,
    )
    db.add(row)
    db.commit()
    return row


def _heartbeat(client, token):
    return client.post(
        "/api/v1/devices/heartbeat",
        json={"device_id": "pi-07", "app_version": "0.1.0", "queue_len": 0},
        headers={"Authorization": f"Bearer {token}"},
    )


# --------------------------------------------------------------------------- live view + edits
def test_live_view_grid_and_edit_with_reason_writes_audit(
    client, db, admin, faculty, course, section, offering, students, login
):
    session = _session(db, course, section, faculty, ended=False)
    _mark(db, session, students[0], AttendanceStatus.PRESENT)
    token = login(admin.email)

    page = client.get(f"/admin/sessions/{session.id}")
    assert page.status_code == 200
    assert 'hx-trigger="every 3s"' in page.text and students[2].usn in page.text
    grid = client.get(f"/admin/sessions/{session.id}/grid")
    assert grid.status_code == 200 and "0.80" in grid.text and "Present 1 of 3" in grid.text

    form = client.get(f"/admin/sessions/{session.id}/students/{students[1].id}/edit")
    assert form.status_code == 200 and 'name="reason"' in form.text

    missing_reason = client.post(
        f"/admin/sessions/{session.id}/students/{students[1].id}",
        data={"status": "present", "reason": "  ", "csrf_token": token},
        headers={"HX-Request": "true"},
    )
    assert missing_reason.status_code == 400 and "reason is required" in missing_reason.text
    assert db.scalar(select(Attendance).where(Attendance.student_id == students[1].id)) is None

    ok = client.post(
        f"/admin/sessions/{session.id}/students/{students[1].id}",
        data={"status": "late", "reason": "arrived at 09:20, bus strike", "csrf_token": token},
        headers={"HX-Request": "true"},
    )
    assert ok.status_code == 200 and "Present 2 of 3" in ok.text
    record = db.scalar(select(Attendance).where(Attendance.student_id == students[1].id))
    assert record is not None
    assert record.status is AttendanceStatus.LATE and record.method is AttendanceMethod.MANUAL
    log = db.scalar(select(AuditLog).where(AuditLog.action == "attendance.edit"))
    assert log is not None and log.actor_id == admin.id
    assert log.reason == "arrived at 09:20, bus strike"
    assert log.before is None and log.after["status"] == "late"
    assert log.after["usn"] == students[1].usn

    # Changing an existing face record keeps the before snapshot and switches to manual.
    changed = client.post(
        f"/admin/sessions/{session.id}/students/{students[0].id}",
        data={
            "status": "absent",
            "reason": "proxy suspected, verified with rep",
            "csrf_token": token,
        },
    )
    assert changed.status_code == 303
    first = db.scalar(select(Attendance).where(Attendance.student_id == students[0].id))
    db.refresh(first)
    assert first.status is AttendanceStatus.ABSENT and first.method is AttendanceMethod.MANUAL
    logs = list(
        db.scalars(
            select(AuditLog).where(AuditLog.action == "attendance.edit").order_by(AuditLog.id)
        )
    )
    assert logs[-1].before["status"] == "present" and logs[-1].before["method"] == "face"
    assert logs[-1].before["score"] == 0.8


def test_bulk_edit_requires_reason_and_audits_each_record(
    client, db, admin, faculty, course, section, offering, students, login
):
    session = _session(db, course, section, faculty)
    token = login(admin.email)
    ids = [str(students[0].id), str(students[1].id)]
    no_reason = client.post(
        f"/admin/sessions/{session.id}/bulk",
        data={"student_id": ids, "status": "excused", "reason": "", "csrf_token": token},
    )
    assert no_reason.status_code == 303
    assert db.scalar(select(Attendance)) is None

    ok = client.post(
        f"/admin/sessions/{session.id}/bulk",
        data={
            "student_id": ids,
            "status": "excused",
            "reason": "college event: tech fest",
            "csrf_token": token,
        },
    )
    assert ok.status_code == 303
    rows = list(db.scalars(select(Attendance).where(Attendance.session_id == session.id)))
    assert len(rows) == 2 and all(r.status is AttendanceStatus.EXCUSED for r in rows)
    logs = list(db.scalars(select(AuditLog).where(AuditLog.action == "attendance.bulk_edit")))
    assert len(logs) == 2 and all(log.reason == "college event: tech fest" for log in logs)


def test_faculty_only_edits_own_courses_within_the_window(
    client, db, admin, faculty, make_faculty, course, section, offering, students, login
):
    other_faculty = make_faculty(email="other@test.local")
    other_course = Course(code="22EC99", name="Other", department="ECE", semester=7)
    db.add(other_course)
    db.flush()
    db.add(
        CourseSection(course_id=other_course.id, section_id=section.id, faculty_id=other_faculty.id)
    )
    db.commit()
    own = _session(db, course, section, faculty)
    foreign = _session(db, other_course, section, other_faculty)
    old = _session(db, course, section, faculty, started_at=datetime.now(UTC) - timedelta(days=10))
    token = login(faculty.email)

    listing = client.get("/admin/sessions")
    assert str(own.id) in listing.text and str(foreign.id) not in listing.text

    assert (
        client.get(f"/admin/sessions/{foreign.id}", headers={"accept": "text/html"}).status_code
        == 403
    )
    denied = client.post(
        f"/admin/sessions/{foreign.id}/students/{students[0].id}",
        data={"status": "present", "reason": "x", "csrf_token": token},
    )
    assert denied.status_code == 403

    allowed = client.post(
        f"/admin/sessions/{own.id}/students/{students[0].id}",
        data={"status": "present", "reason": "was in the lab", "csrf_token": token},
    )
    assert allowed.status_code == 303

    expired_page = client.get(f"/admin/sessions/{old.id}")
    assert expired_page.status_code == 200 and "edit window" in expired_page.text
    expired = client.post(
        f"/admin/sessions/{old.id}/students/{students[0].id}",
        data={"status": "present", "reason": "late request", "csrf_token": token},
    )
    assert expired.status_code == 403

    # Admin is not bound by the window.
    client.post("/admin/logout", data={"csrf_token": token})
    admin_token = login(admin.email)
    approved = client.post(
        f"/admin/sessions/{old.id}/students/{students[0].id}",
        data={"status": "present", "reason": "approved by HoD", "csrf_token": admin_token},
    )
    assert approved.status_code == 303


def test_end_session_from_browser_materialises_absents(
    client, db, admin, faculty, course, section, offering, students, login
):
    session = _session(db, course, section, faculty, ended=False)
    _mark(db, session, students[0], AttendanceStatus.PRESENT)
    token = login(admin.email)
    response = client.post(f"/admin/sessions/{session.id}/end", data={"csrf_token": token})
    assert response.status_code == 303
    db.refresh(session)
    assert session.status is SessionStatus.ENDED
    statuses = sorted(
        s.value
        for s in db.scalars(select(Attendance.status).where(Attendance.session_id == session.id))
    )
    assert statuses == ["absent", "absent", "present"]


# --------------------------------------------------------------------------- reports
def test_reports_percentages_shortage_and_exports(
    client, db, admin, faculty, course, section, offering, students, login
):
    s1 = _session(db, course, section, faculty, started_at=datetime.now(UTC) - timedelta(days=2))
    s2 = _session(db, course, section, faculty, started_at=datetime.now(UTC) - timedelta(days=1))
    a, b, c = students
    _mark(db, s1, a, AttendanceStatus.PRESENT)
    _mark(db, s2, a, AttendanceStatus.LATE)
    _mark(db, s1, b, AttendanceStatus.PRESENT)
    _mark(db, s2, b, AttendanceStatus.EXCUSED, method=AttendanceMethod.MANUAL)
    _mark(db, s1, c, AttendanceStatus.ABSENT, method=AttendanceMethod.MANUAL)
    _mark(db, s2, c, AttendanceStatus.ABSENT, method=AttendanceMethod.MANUAL)
    login(admin.email)

    page = client.get("/admin/reports")
    assert page.status_code == 200
    by_usn = {match[1]: match for match in ROW_RE.findall(page.text)}
    assert by_usn[a.usn][2:] == ("2", "2", "1", "0", "100%")
    assert by_usn[b.usn][2:] == ("2", "1", "0", "1", "100%")  # excused leaves the denominator
    assert by_usn[c.usn][2:] == ("2", "0", "0", "0", "0%") and by_usn[c.usn][0] == "short"

    shortage = client.get("/admin/reports", params={"view": "shortage"})
    assert c.usn in shortage.text and a.usn not in shortage.text

    csv_export = client.get("/admin/reports/export", params={"fmt": "csv"})
    assert csv_export.status_code == 200
    assert csv_export.headers["content-type"].startswith("text/csv")
    assert f"22EC71,VLSI Design,{a.usn}" in csv_export.text and "100.0" in csv_export.text

    xlsx_export = client.get("/admin/reports/export", params={"fmt": "xlsx", "view": "shortage"})
    assert xlsx_export.status_code == 200
    sheet = load_workbook(io.BytesIO(xlsx_export.content)).active
    values = [[cell.value for cell in row] for row in sheet.iter_rows()]
    assert values[0][:4] == ["course_code", "course_name", "usn", "name"]
    assert len(values) == 2 and values[1][2] == c.usn


def test_dashboard_shows_devices_warnings_and_shortage(
    client, db, admin, faculty, course, section, offering, students, device, login
):
    row, _ = device
    row.last_seen_at = datetime.now(UTC) - timedelta(minutes=30)
    row.queue_len = 4
    row.clock_synced = False
    db.commit()
    s1 = _session(db, course, section, faculty, device=row)
    _mark(db, s1, students[0], AttendanceStatus.PRESENT)
    login(admin.email)
    page = client.get("/admin")
    assert page.status_code == 200
    assert "offline for more than 10 minutes" in page.text
    assert "4 records waiting to sync" in page.text and "clock not NTP-synced" in page.text
    assert students[1].usn in page.text  # below threshold (0%)
    assert "Sessions today" in page.text and "22EC71" in page.text


# --------------------------------------------------------------------------- devices
def test_every_post_form_carries_the_csrf_token(client, admin, section, login):
    """A browser only sends what the page contains; posting the token by hand hid a
    register form that had none (every submit was a 403)."""
    token = login(admin.email)
    client.post("/admin/devices", data={"name": "pi-08", "csrf_token": token})
    for path in (
        "/admin/devices",
        "/admin/students",
        "/admin/students/new",
        "/admin/courses",
        "/admin/sections",
        "/admin/periods",
        "/admin/faculty",
        "/admin/me",
        "/admin/enrolment",
    ):
        page = client.get(path)
        assert page.status_code == 200, path
        for attrs, body in FORM_RE.findall(page.text):
            if 'method="post"' in attrs.lower():
                assert f'name="csrf_token" value="{token}"' in body, f"{path}: <form{attrs}>"


def test_device_register_assign_rotate_revoke(client, db, admin, section, login):
    token = login(admin.email)
    created = client.post("/admin/devices", data={"name": "pi-07", "csrf_token": token})
    assert created.status_code == 200
    device_token = TOKEN_RE.search(created.text).group(1)
    assert _heartbeat(client, device_token).status_code == 200
    device = db.scalar(select(Device).where(Device.name == "pi-07"))
    assert device is not None

    assigned = client.post(
        f"/admin/devices/{device.id}/sections",
        data={"section_id": [str(section.id)], "csrf_token": token},
    )
    assert assigned.status_code == 303
    assert _heartbeat(client, device_token).json()["assigned_section_ids"] == [section.id]

    listing = client.get("/admin/devices")
    assert "pi-07" in listing.text and "online" in listing.text

    revoked = client.post(
        f"/admin/devices/{device.id}/revoke", data={"reason": "lost in lab", "csrf_token": token}
    )
    assert revoked.status_code == 303
    denied = _heartbeat(client, device_token)
    assert denied.status_code == 401 and denied.json()["detail"] == "Device revoked."

    rotated = client.post(f"/admin/devices/{device.id}/rotate", data={"csrf_token": token})
    new_token = TOKEN_RE.search(rotated.text).group(1)
    assert new_token != device_token
    assert _heartbeat(client, new_token).status_code == 200
    actions = set(db.scalars(select(AuditLog.action).where(AuditLog.entity == "device")))
    expected = {"device.register", "device.assign_sections", "device.revoke", "device.rotate_token"}
    assert expected <= actions
    assert client.get("/admin/devices").status_code == 200


def test_faculty_cannot_open_admin_only_pages(client, faculty, login):
    login(faculty.email)
    for path in ("/admin/devices", "/admin/audit", "/admin/enrolment"):
        assert client.get(path, headers={"accept": "text/html"}).status_code == 403


# --------------------------------------------------------------------------- student page
def test_student_page_history_reenrol_and_withdraw_consent(
    client, db, admin, faculty, course, section, offering, students, settings, login
):
    student = students[0]
    consent = Consent(
        student_id=student.id,
        consent_at=datetime.now(UTC),
        consent_version="v1",
        method=ConsentMethod.PAPER,
    )
    db.add(consent)
    db.flush()

    def template():
        return FaceTemplate(
            student_id=student.id,
            consent_id=consent.id,
            embedding=b"\0" * EMBEDDING_BYTES,
            model_version=settings.model_version,
            source=TemplateSource.DEVICE,
        )

    db.add_all([template(), template()])
    db.commit()
    s1 = _session(db, course, section, faculty)
    _mark(db, s1, student, AttendanceStatus.PRESENT)
    token = login(admin.email)

    page = client.get(f"/admin/students/{student.id}")
    assert page.status_code == 200
    assert "Face templates (2)" in page.text and "22EC71" in page.text and "100%" in page.text
    assert "Withdraw consent" in page.text

    no_reason = client.post(
        f"/admin/students/{student.id}/reenrol", data={"reason": "", "csrf_token": token}
    )
    assert no_reason.status_code == 303
    assert db.scalar(select(FaceTemplate).where(FaceTemplate.student_id == student.id)) is not None

    reenrol = client.post(
        f"/admin/students/{student.id}/reenrol",
        data={"reason": "poor templates", "csrf_token": token},
    )
    assert reenrol.status_code == 303
    assert db.scalar(select(FaceTemplate).where(FaceTemplate.student_id == student.id)) is None
    db.refresh(consent)
    assert consent.withdrawn_at is None

    db.add(template())
    db.commit()
    withdraw = client.post(
        f"/admin/students/{student.id}/withdraw-consent",
        data={"reason": "student request", "csrf_token": token},
    )
    assert withdraw.status_code == 303
    db.refresh(consent)
    assert consent.withdrawn_at is not None
    assert db.scalar(select(FaceTemplate).where(FaceTemplate.student_id == student.id)) is None
    assert db.scalar(select(Attendance).where(Attendance.student_id == student.id)) is not None
    withdrawal = db.scalar(select(AuditLog).where(AuditLog.action == "consent.withdraw"))
    assert withdrawal is not None and withdrawal.reason == "student request"
    after = client.get(f"/admin/students/{student.id}")
    assert "withdrawn" in after.text and "Not enrolled yet" in after.text


# --------------------------------------------------------------------------- audit + enrolment
def test_audit_page_filters(client, db, admin, faculty, students, login):
    token = login(admin.email)
    client.post(
        "/admin/students",
        data={
            "usn": "1BM22EC077",
            "name": "Audit Me",
            "department": "ECE",
            "semester": "7",
            "csrf_token": token,
        },
    )
    page = client.get("/admin/audit")
    assert page.status_code == 200 and "student.create" in page.text and "auth.login" in page.text
    filtered = client.get("/admin/audit", params={"action": "student.", "actor": admin.id})
    assert "student.create" in filtered.text and "auth.login" not in filtered.text
    nothing = client.get("/admin/audit", params={"entity": "device"})
    assert "Nothing matches" in nothing.text


def test_enrolment_page_lists_missing_consent_and_templates(
    client, db, admin, students, settings, login
):
    consent = Consent(
        student_id=students[0].id,
        consent_at=datetime.now(UTC),
        consent_version="v1",
        method=ConsentMethod.PAPER,
    )
    db.add(consent)
    db.commit()
    login(admin.email)
    page = client.get("/admin/enrolment")
    assert page.status_code == 200
    assert f"Missing consent ({len(students) - 1})" in page.text
    assert "Consented but not enrolled (1)" in page.text and students[0].usn in page.text
