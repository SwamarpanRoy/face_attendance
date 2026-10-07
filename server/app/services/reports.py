"""Attendance percentages, shortage lists and exports.

Counting rule (documented on the reports page): a student's percentage for a course is
``attended / (held - excused)`` where *held* is the number of ended sessions for the
student's section in the date range, *attended* counts ``present`` and ``late``, and
``excused`` sessions are removed from the denominator. Sessions still live are not
counted.
"""

from __future__ import annotations

import csv
import io
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from server.app.models import (
    Attendance,
    AttendanceSession,
    AttendanceStatus,
    Course,
    SessionStatus,
    Student,
)

ATTENDED = (AttendanceStatus.PRESENT, AttendanceStatus.LATE)


@dataclass
class StudentCourseStat:
    student_id: int
    usn: str
    name: str
    section_id: int | None
    course_id: int
    course_code: str
    course_name: str
    held: int = 0
    attended: int = 0
    excused: int = 0
    late: int = 0

    @property
    def counted(self) -> int:
        return max(0, self.held - self.excused)

    @property
    def percentage(self) -> float | None:
        return None if self.counted == 0 else 100.0 * self.attended / self.counted


@dataclass
class HistoryRow:
    session: AttendanceSession
    record: Attendance | None


@dataclass
class StudentHistory:
    per_course: list[StudentCourseStat]
    rows: list[HistoryRow] = field(default_factory=list)


def date_bounds(
    start: date | None, end: date | None, tz: str
) -> tuple[datetime | None, datetime | None]:
    """Inclusive local calendar days -> aware UTC bounds for ``started_at`` filters."""
    zone = ZoneInfo(tz)
    lower = datetime.combine(start, time.min, tzinfo=zone) if start else None
    upper = datetime.combine(end + timedelta(days=1), time.min, tzinfo=zone) if end else None
    return lower, upper


def _ended_sessions(
    db: Session,
    *,
    course_ids: set[int] | None,
    section_id: int | None,
    lower: datetime | None,
    upper: datetime | None,
) -> list[AttendanceSession]:
    stmt = (
        select(AttendanceSession)
        .options(selectinload(AttendanceSession.course))
        .where(AttendanceSession.status == SessionStatus.ENDED)
    )
    if course_ids is not None:
        stmt = stmt.where(AttendanceSession.course_id.in_(course_ids))
    if section_id is not None:
        stmt = stmt.where(AttendanceSession.section_id == section_id)
    if lower is not None:
        stmt = stmt.where(AttendanceSession.started_at >= lower)
    if upper is not None:
        stmt = stmt.where(AttendanceSession.started_at < upper)
    return list(db.scalars(stmt.order_by(AttendanceSession.started_at)))


def course_matrix(
    db: Session,
    *,
    course_ids: set[int] | None,
    section_id: int | None,
    lower: datetime | None,
    upper: datetime | None,
) -> list[StudentCourseStat]:
    """One row per (student, course) over the ended sessions in range."""
    sessions = _ended_sessions(
        db, course_ids=course_ids, section_id=section_id, lower=lower, upper=upper
    )
    if not sessions:
        return []
    session_ids = [s.id for s in sessions]
    by_section_course: dict[tuple[int, int], int] = defaultdict(int)
    for s in sessions:
        by_section_course[(s.section_id, s.course_id)] += 1
    courses = {s.course_id: s.course for s in sessions}
    section_ids = {s.section_id for s in sessions}
    students = list(
        db.scalars(
            select(Student).where(Student.section_id.in_(section_ids), Student.active.is_(True))
        )
    )
    records = db.execute(
        select(Attendance.student_id, AttendanceSession.course_id, Attendance.status)
        .join(AttendanceSession, AttendanceSession.id == Attendance.session_id)
        .where(Attendance.session_id.in_(session_ids))
    ).all()

    stats: dict[tuple[int, int], StudentCourseStat] = {}
    for student in students:
        for (section_id_, course_id), held in by_section_course.items():
            if section_id_ != student.section_id:
                continue
            course: Course = courses[course_id]
            stats[(student.id, course_id)] = StudentCourseStat(
                student_id=student.id,
                usn=student.usn,
                name=student.name,
                section_id=student.section_id,
                course_id=course_id,
                course_code=course.code,
                course_name=course.name,
                held=held,
            )
    for student_id, course_id, status in records:
        stat = stats.get((int(student_id), int(course_id)))
        if stat is None:
            continue
        if status in ATTENDED:
            stat.attended += 1
        if status is AttendanceStatus.LATE:
            stat.late += 1
        if status is AttendanceStatus.EXCUSED:
            stat.excused += 1
    return sorted(stats.values(), key=lambda s: (s.course_code, s.usn))


def shortage(rows: list[StudentCourseStat], threshold_pct: float) -> list[StudentCourseStat]:
    return sorted(
        (r for r in rows if r.percentage is not None and r.percentage < threshold_pct),
        key=lambda r: (r.percentage or 0.0, r.course_code, r.usn),
    )


def student_history(db: Session, student: Student, *, tz: str) -> StudentHistory:
    """Every session of the student's section (ended or live) with the student's record."""
    sessions = list(
        db.scalars(
            select(AttendanceSession)
            .options(
                selectinload(AttendanceSession.course), selectinload(AttendanceSession.faculty)
            )
            .where(AttendanceSession.section_id == student.section_id)
            .order_by(AttendanceSession.started_at.desc())
        )
    )
    records = {
        r.session_id: r
        for r in db.scalars(select(Attendance).where(Attendance.student_id == student.id))
    }
    per_course = [
        s
        for s in course_matrix(
            db, course_ids=None, section_id=student.section_id, lower=None, upper=None
        )
        if s.student_id == student.id
    ]
    return StudentHistory(
        per_course=per_course, rows=[HistoryRow(s, records.get(s.id)) for s in sessions]
    )


# --------------------------------------------------------------------------- exports
HEADERS = [
    "course_code",
    "course_name",
    "usn",
    "name",
    "held",
    "attended",
    "late",
    "excused",
    "percentage",
]


def _row_values(stat: StudentCourseStat) -> list[Any]:
    pct = stat.percentage
    return [
        stat.course_code,
        stat.course_name,
        stat.usn,
        stat.name,
        stat.held,
        stat.attended,
        stat.late,
        stat.excused,
        round(pct, 1) if pct is not None else "",
    ]


def export_csv(rows: list[StudentCourseStat]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(HEADERS)
    for stat in rows:
        writer.writerow(_row_values(stat))
    return buffer.getvalue().encode("utf-8-sig")


def export_xlsx(rows: list[StudentCourseStat], *, title: str = "Attendance") -> bytes:
    workbook = Workbook()
    sheet = workbook.create_sheet(title=title[:31], index=0)
    if "Sheet" in workbook.sheetnames:
        del workbook["Sheet"]
    table = [HEADERS, *[_row_values(stat) for stat in rows]]
    for values in table:
        sheet.append(values)
    for index, _header in enumerate(HEADERS, start=1):
        width = max(len(str(row[index - 1])) for row in table)
        sheet.column_dimensions[get_column_letter(index)].width = min(40, max(10, width + 2))
    sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
